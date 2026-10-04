"""
MK DATA — Enterprise Client Handover Document Generator
Generates high-resolution, professionally styled .docx and .pdf documents
for offline review and client repository delivery.
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
    KeepTogether,
)

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "docs")
os.makedirs(OUTPUT_DIR, exist_ok=True)

DOCX_PATH = os.path.join(OUTPUT_DIR, "MK_DATA_ENTERPRISE_HANDOVER_GUIDE.docx")
PDF_PATH = os.path.join(OUTPUT_DIR, "MK_DATA_ENTERPRISE_HANDOVER_GUIDE.pdf")


def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def generate_docx():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    title_style = doc.styles.add_style("CoverTitle", 1)
    title_style.font.name = "Calibri"
    title_style.font.size = Pt(26)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(15, 23, 42)  # Slate 900

    # Title Banner
    p_title = doc.add_paragraph("MK DATA 📱", style="CoverTitle")
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_sub = doc.add_paragraph("Enterprise Fintech & VTU Platform — Handover & Operations Guide")
    p_sub.runs[0].font.size = Pt(14)
    p_sub.runs[0].font.color.rgb = RGBColor(71, 85, 105)
    p_sub.runs[0].font.bold = True
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_meta = doc.add_paragraph(
        "Client Handover Deliverable | Prepared for Production Deployment\nLead Developer Contact: 08034910470"
    )
    p_meta.runs[0].font.size = Pt(10)
    p_meta.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 1: Executive Overview
    h1 = doc.add_heading("1. Executive System Overview", level=1)
    h1.runs[0].font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph(
        "MK DATA is an enterprise-grade fintech and VTU (Virtual Top-Up) platform engineered to deliver "
        "instant mobile data bundles, airtime recharge, electricity utility tokens, cable TV subscriptions, "
        "and educational exam PINs across Nigeria. The system is built on Next.js 16 App Router, React 19, "
        "PostgreSQL (Neon Serverless), Prisma ORM, and TailwindCSS, backed by atomic wallet debit safeguards "
        "and BillStack automated virtual bank account funding."
    )

    # Section 2: Technology Stack & Core Components
    h2 = doc.add_heading("2. Core Technology Stack", level=1)
    h2.runs[0].font.color.rgb = RGBColor(30, 41, 59)

    tech_table = doc.add_table(rows=1, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tech_table.rows[0].cells
    hdr_cells[0].text = "Component Layer"
    hdr_cells[1].text = "Technology / Framework"
    hdr_cells[2].text = "Purpose / Role"

    for cell in hdr_cells:
        set_cell_background(cell, "1E293B")
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.bold = True
                r.font.size = Pt(9.5)

    tech_data = [
        ("Web Frontend & API", "Next.js 16 (App Router), React 19, TypeScript", "Responsive client portal, checkout flows & API routes"),
        ("Database Layer", "PostgreSQL (Neon Serverless) & Prisma 6", "ACID financial transactions, relational catalogs & indexing"),
        ("Payment Processing", "BillStack Dedicated Virtual Accounts", "Instant wallet funding via Wema/Providus bank transfers"),
        ("Telecom Providers", "Alrahuz Data, AmySub, Saiful, SMEPlug", "Multi-vendor mobile data, airtime, and utility fulfillment"),
        ("Push Notifications", "Firebase Admin SDK (FCM Multicast)", "Automated device push notifications and broadcast alerts"),
        ("Authentication", "Jose JWT, Bcrypt & Biometric Tokens", "Session management, PIN hashing & biometric device auth"),
    ]

    for cat, tech, purpose in tech_data:
        row = tech_table.add_row().cells
        row[0].text = cat
        row[1].text = tech
        row[2].text = purpose
        for i, cell in enumerate(row):
            set_cell_background(cell, "F8FAFC" if len(tech_table.rows) % 2 == 0 else "FFFFFF")
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 3: Third-Party Account Transfer Matrix
    h3 = doc.add_heading("3. Third-Party Account Transfer Matrix", level=1)
    h3.runs[0].font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph(
        "To assume complete operational ownership of MK DATA, transfer the following services to client-owned accounts:"
    )

    transfer_table = doc.add_table(rows=1, cols=3)
    transfer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_hdrs = transfer_table.rows[0].cells
    t_hdrs[0].text = "Service Provider"
    t_hdrs[1].text = "Account Role"
    t_hdrs[2].text = "Transfer & Verification Action"

    for cell in t_hdrs:
        set_cell_background(cell, "0F766E")  # Teal 700
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.bold = True
                r.font.size = Pt(9.5)

    transfer_data = [
        ("Vercel Hosting", "Web & API Host", "Transfer Team / Project to client Vercel organization. Ensure production domain SSL is active."),
        ("Neon PostgreSQL", "Database Host", "Transfer project ownership in Neon Console. Verify database connectivity."),
        ("BillStack Payments", "Payment Gateway", "Update KYC, settlement bank account, and Webhook URL to production domain."),
        ("Alrahuz Data", "Primary Telecom Vendor", "Register/transfer account, fund vendor wallet float, and update ALRAHUZ_API_TOKEN."),
        ("Firebase (Google Cloud)", "Push Notifications", "Grant Owner role in Firebase Console. Generate and upload fresh service account JSON."),
        ("Domain Registrar", "DNS & Domain", "Point DNS A/CNAME records to Vercel and verify https://mkdata.com.ng."),
    ]

    for srv, role, action in transfer_data:
        row = transfer_table.add_row().cells
        row[0].text = srv
        row[1].text = role
        row[2].text = action
        for cell in row:
            set_cell_background(cell, "F0FDFA" if len(transfer_table.rows) % 2 == 0 else "FFFFFF")
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 4: Operational Troubleshooting & FAQs
    h4 = doc.add_heading("4. Operational Troubleshooting Matrix", level=1)
    h4.runs[0].font.color.rgb = RGBColor(30, 41, 59)

    diag_table = doc.add_table(rows=1, cols=3)
    diag_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    d_hdrs = diag_table.rows[0].cells
    d_hdrs[0].text = "Issue Symptom"
    d_hdrs[1].text = "Root Cause"
    d_hdrs[2].text = "Resolution Action"

    for cell in d_hdrs:
        set_cell_background(cell, "334155")
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.bold = True
                r.font.size = Pt(9.5)

    diag_data = [
        ("Bank transfer made, wallet not credited", "Webhook signature mismatch or URL misconfigured", "Verify BILLSTACK_SECRET_KEY in Vercel and check Admin > Webhooks log."),
        ("Data order fails: Vendor balance low", "Alrahuz float wallet ran out of funds", "Immediately top up Alrahuz account via bank transfer. Orders can be retried."),
        ("Push notification broadcast error", "FIREBASE_SERVICE_ACCOUNT_JSON missing/invalid", "Re-export service account JSON from Google Cloud Console and paste in Vercel."),
        ("Admin login PIN rejected", "Admin user PIN reset or missing in DB", "Run 'npm run seed:admin-services' with MK_ADMIN_PIN='000000'."),
    ]

    for sym, cause, resol in diag_data:
        row = diag_table.add_row().cells
        row[0].text = sym
        row[1].text = cause
        row[2].text = resol
        for cell in row:
            set_cell_background(cell, "F8FAFC" if len(diag_table.rows) % 2 == 0 else "FFFFFF")
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Developer Contact Banner
    h5 = doc.add_heading("5. Lead Developer Contact & Technical Handover Support", level=1)
    h5.runs[0].font.color.rgb = RGBColor(30, 41, 59)

    p_contact = doc.add_paragraph()
    p_contact.add_run("For source code walkthroughs, deployment inquiries, or technical assistance:\n").bold = True
    p_contact.add_run("• Phone / WhatsApp: ").bold = False
    p_contact.add_run("08034910470\n").bold = True
    p_contact.add_run("• Handover Support: Direct Architecture & Environment Inquiries\n")

    doc.save(DOCX_PATH)
    print(f"Generated DOCX: {DOCX_PATH}")


def generate_pdf():
    pdf_doc = SimpleDocTemplate(
        PDF_PATH,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "CoverTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=26,
        textColor=colors.HexColor("#0F172A"),
        alignment=1,  # Center
    )

    sub_style = ParagraphStyle(
        "CoverSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#0F766E"),
        alignment=1,
    )

    meta_style = ParagraphStyle(
        "CoverMeta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#64748B"),
        alignment=1,
    )

    h1_style = ParagraphStyle(
        "Heading1Style",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=17,
        textColor=colors.HexColor("#1E293B"),
        spaceBefore=12,
        spaceAfter=6,
    )

    body_style = ParagraphStyle(
        "BodyStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#334155"),
        spaceAfter=8,
    )

    table_cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#1E293B"),
    )

    table_hdr_style = ParagraphStyle(
        "TableHdr",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
    )

    story = []

    # Title Banner
    story.append(Paragraph("MK DATA 📱", title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph("Enterprise Fintech & VTU Platform — Client Handover Guide", sub_style))
    story.append(Spacer(1, 4))
    story.append(
        Paragraph(
            "Official Production Handover Deliverable | Lead Developer: <b>08034910470</b>",
            meta_style,
        )
    )
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#0F766E"), spaceAfter=12))

    # 1. Executive Summary
    story.append(Paragraph("1. Executive Summary & Architecture", h1_style))
    story.append(
        Paragraph(
            "MK DATA is an enterprise-grade fintech and VTU (Virtual Top-Up) solution built on Next.js 16 "
            "App Router, React 19, Neon Serverless PostgreSQL, Prisma ORM, and TailwindCSS. The platform supports "
            "instant mobile data bundling, airtime recharging, electricity token vending, cable TV subscriptions, "
            "and WAEC/NECO/JAMB exam PIN purchases. Automated wallet funding is achieved via dedicated BillStack "
            "reserved virtual bank accounts with HMAC webhook validation.",
            body_style,
        )
    )

    # 2. Technology Stack Table
    story.append(Paragraph("2. Technology Stack & Frameworks", h1_style))
    tech_rows = [
        [
            Paragraph("<b>Component Layer</b>", table_hdr_style),
            Paragraph("<b>Technology / Framework</b>", table_hdr_style),
            Paragraph("<b>Operational Role</b>", table_hdr_style),
        ],
        [
            Paragraph("Frontend & API", table_cell_style),
            Paragraph("Next.js 16, React 19, TypeScript", table_cell_style),
            Paragraph("Web portal, PWA, order checkout & API endpoints", table_cell_style),
        ],
        [
            Paragraph("Database", table_cell_style),
            Paragraph("PostgreSQL (Neon Serverless) & Prisma 6", table_cell_style),
            Paragraph("ACID transactions, indexed catalogs & ledgers", table_cell_style),
        ],
        [
            Paragraph("Payment Gateway", table_cell_style),
            Paragraph("BillStack Reserved Accounts", table_cell_style),
            Paragraph("Automated bank transfer wallet funding", table_cell_style),
        ],
        [
            Paragraph("VTU Telecom APIs", table_cell_style),
            Paragraph("Alrahuz Data, AmySub, Saiful, SMEPlug", table_cell_style),
            Paragraph("Multi-network data & utility dispatch engine", table_cell_style),
        ],
        [
            Paragraph("Push Notifications", table_cell_style),
            Paragraph("Firebase Admin SDK (FCM Multicast)", table_cell_style),
            Paragraph("Device push notifications & broadcast notices", table_cell_style),
        ],
    ]

    t1 = Table(tech_rows, colWidths=[110, 190, 230])
    t1.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(t1)
    story.append(Spacer(1, 10))

    # 3. Third-Party Account Transfer Matrix
    story.append(Paragraph("3. Third-Party Account Transfer Matrix", h1_style))
    transfer_rows = [
        [
            Paragraph("<b>Provider</b>", table_hdr_style),
            Paragraph("<b>Architecture Role</b>", table_hdr_style),
            Paragraph("<b>Required Transfer & Verification Action</b>", table_hdr_style),
        ],
        [
            Paragraph("Vercel Hosting", table_cell_style),
            Paragraph("Web & API Hosting", table_cell_style),
            Paragraph("Transfer Project ownership; verify custom domain SSL.", table_cell_style),
        ],
        [
            Paragraph("Neon PostgreSQL", table_cell_style),
            Paragraph("Production Database", table_cell_style),
            Paragraph("Transfer project to client organization; verify query execution.", table_cell_style),
        ],
        [
            Paragraph("BillStack", table_cell_style),
            Paragraph("Payment Gateway", table_cell_style),
            Paragraph("Update KYC, settlement bank, and Webhook URL to live domain.", table_cell_style),
        ],
        [
            Paragraph("Alrahuz Data", table_cell_style),
            Paragraph("Primary Telecom API", table_cell_style),
            Paragraph("Transfer portal account, fund float balance, and set API token.", table_cell_style),
        ],
        [
            Paragraph("Firebase (GCP)", table_cell_style),
            Paragraph("Push Notifications", table_cell_style),
            Paragraph("Grant Owner role in Firebase Console; generate fresh JSON key.", table_cell_style),
        ],
    ]

    t2 = Table(transfer_rows, colWidths=[110, 140, 280])
    t2.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F766E")),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0FDFA")]),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(t2)
    story.append(Spacer(1, 10))

    # 4. Troubleshooting Matrix
    story.append(Paragraph("4. Operational Troubleshooting Matrix", h1_style))
    diag_rows = [
        [
            Paragraph("<b>Symptom</b>", table_hdr_style),
            Paragraph("<b>Probable Cause</b>", table_hdr_style),
            Paragraph("<b>Resolution Procedure</b>", table_hdr_style),
        ],
        [
            Paragraph("Bank transfer not credited", table_cell_style),
            Paragraph("Webhook URL error or secret mismatch", table_cell_style),
            Paragraph("Verify BILLSTACK_SECRET_KEY in Vercel; check /admin/webhooks log.", table_cell_style),
        ],
        [
            Paragraph("Data vending failed", table_cell_style),
            Paragraph("Vendor float balance exhausted", table_cell_style),
            Paragraph("Top up Alrahuz vendor balance via bank transfer immediately.", table_cell_style),
        ],
        [
            Paragraph("Push broadcast fails", table_cell_style),
            Paragraph("Firebase service JSON invalid", table_cell_style),
            Paragraph("Re-generate private key from Google Cloud Console & update Vercel.", table_cell_style),
        ],
        [
            Paragraph("Admin login rejected", table_cell_style),
            Paragraph("PIN mismatch or unseeded user", table_cell_style),
            Paragraph("Run 'npm run seed:admin-services' with MK_ADMIN_PIN='000000'.", table_cell_style),
        ],
    ]

    t3 = Table(diag_rows, colWidths=[130, 160, 240])
    t3.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(t3)
    story.append(Spacer(1, 12))

    # 5. Developer Contact Box
    contact_box = [
        [
            Paragraph(
                "<b>📞 Lead Developer Contact & Technical Handover Support</b><br/>"
                "• <b>Phone / WhatsApp</b>: <b>08034910470</b><br/>"
                "• <b>Handover Assistance</b>: Direct Architecture, Database & Environment Walkthroughs",
                ParagraphStyle(
                    "ContactBox",
                    parent=styles["Normal"],
                    fontName="Helvetica",
                    fontSize=9,
                    leading=14,
                    textColor=colors.HexColor("#0F172A"),
                ),
            )
        ]
    ]
    t_contact = Table(contact_box, colWidths=[530])
    t_contact.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#E2E8F0")),
                ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#0F766E")),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 12),
                ("RIGHTPADDING", (0, 0), (-1, -1), 12),
            ]
        )
    )
    story.append(t_contact)

    pdf_doc.build(story)
    print(f"Generated PDF: {PDF_PATH}")


if __name__ == "__main__":
    generate_docx()
    generate_pdf()
